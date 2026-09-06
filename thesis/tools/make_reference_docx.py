#!/usr/bin/env python3
"""
Generate thesis/reference.docx, the Pandoc style template that makes the
.docx output land in University of Buea format instead of Pandoc defaults.

Applies the measured spec from report/UB-PhD-Thesis-Format-Spec.md:

    Normal      Times New Roman 12 pt, double line spacing, justified,
                no space before/after, 1.27 cm first-line indent
    Heading 1   Times New Roman 16 pt bold, all caps, centred
    Heading 2   Times New Roman 14 pt bold
    Heading 3   Times New Roman 14 pt bold (same as level 2 in the exemplar)
    Margins     left 3.5 cm, top/right/bottom 2.5 cm
    Page number top right (added as a field in the header)

Run once:  python3 tools/make_reference_docx.py
Then build.sh picks reference.docx up automatically.
"""

import subprocess
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

HERE = Path(__file__).resolve().parent.parent
REF = HERE / "reference.docx"

FONT = "Times New Roman"


def get_style(doc, name):
    """Look up by display name, falling back to style_id.

    Pandoc's default reference.docx stores heading styles with ids like
    'Heading1'. python-docx's name lookup does not always find them, and its
    style_id fallback is deprecated, so resolve explicitly.
    """
    for s in doc.styles:
        if s.name == name or s.style_id == name.replace(" ", ""):
            return s
    raise KeyError(name)


def set_font(style, size_pt, bold=False, caps=False, black=True):
    """Pin a style to Times New Roman at a given size.

    Two traps in Pandoc's default reference.docx:

      1. Heading styles reference *theme* fonts via w:rFonts/@w:asciiTheme.
         Word resolves the theme in preference to @w:ascii, so setting the
         font name alone leaves headings in Calibri. The theme attributes
         have to be removed.
      2. Heading styles carry an accent colour (4F81BD, blue). Left alone,
         every heading in the thesis prints blue.
    """
    f = style.font
    f.name = FONT
    f.size = Pt(size_pt)
    f.bold = bold

    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)          # schema: w:rFonts comes first in w:rPr
    for themed in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(themed)) is not None:
            del rfonts.attrib[qn(themed)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)

    if black:
        for col in rpr.findall(qn("w:color")):
            rpr.remove(col)
        col = OxmlElement("w:color")
        col.set(qn("w:val"), "000000")
        rpr.append(col)

    for existing in rpr.findall(qn("w:caps")):
        rpr.remove(existing)
    if caps:
        el = OxmlElement("w:caps")
        el.set(qn("w:val"), "true")
        rpr.append(el)


def add_page_number_header(section):
    """Top-right PAGE field, per guide §3.1f."""
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
        run._r.append(el)
    run.font.name = FONT
    run.font.size = Pt(12)


def main():
    # Start from Pandoc's own default so every style Pandoc emits exists.
    subprocess.run(
        ["pandoc", "-o", str(REF), "--print-default-data-file", "reference.docx"],
        check=True, stdout=subprocess.DEVNULL,
    )
    doc = Document(REF)

    # --- Normal -------------------------------------------------------------
    normal = get_style(doc, "Normal")
    set_font(normal, 12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # No space before or after paragraphs. The exemplar shows every baseline
    # gap at one double-spaced line with no paragraph gap anywhere; paragraphs
    # are separated by a first-line indent instead.
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.first_line_indent = Cm(1.27)

    # --- Headings -----------------------------------------------------------
    for name, size, caps, align in (
        # Sizes measured from the exemplar: chapter 16 pt caps centred,
        # section AND subsection both 14 pt bold, level 4 drops to 12 pt.
        ("Heading 1", 16, True,  WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 14, False, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 14, False, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 4", 12, False, WD_ALIGN_PARAGRAPH.LEFT),
    ):
        try:
            st = get_style(doc, name)
        except KeyError:
            continue
        set_font(st, size, bold=True, caps=caps)
        st.paragraph_format.alignment = align
        st.paragraph_format.space_before = Pt(18)
        st.paragraph_format.space_after = Pt(12)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    # --- Body Text: what Pandoc actually applies to most paragraphs ---------
    # Inheriting from Normal is not enough; Pandoc's Body Text carries its own
    # spacing, so paragraph gaps reappear in Word unless it is set explicitly.
    for name in ("Body Text", "First Paragraph", "Compact"):
        try:
            st = get_style(doc, name)
        except KeyError:
            continue
        set_font(st, 12)
        bpf = st.paragraph_format
        bpf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        bpf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        bpf.space_after = Pt(0)
        bpf.space_before = Pt(0)
        bpf.first_line_indent = Cm(1.27)

    # --- Single-spaced supporting styles ------------------------------------
    for name in ("Caption", "Footnote Text", "Quote", "Bibliography",
                 "Source Code", "Table Caption", "Image Caption"):
        try:
            st = get_style(doc, name)
        except KeyError:
            continue
        set_font(st, 12)
        cpf = st.paragraph_format
        cpf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        cpf.first_line_indent = Cm(0)
        cpf.space_after = Pt(0)
        cpf.space_before = Pt(0)
        if "Caption" in name:
            cpf.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- Page setup ---------------------------------------------------------
    for section in doc.sections:
        section.left_margin = Cm(3.5)      # binding margin, guide §3.1e
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.header_distance = Cm(1.5)  # >= 1.5 cm into the page, §3.1f
        add_page_number_header(section)

    doc.save(REF)
    print(f"wrote {REF.relative_to(HERE.parent)}")


if __name__ == "__main__":
    main()
