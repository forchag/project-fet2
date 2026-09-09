#!/usr/bin/env python3
"""Create the editable cover letter and highlights files for submission."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


HERE = Path(__file__).resolve().parent
FONT = "Times New Roman"


def configure(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    style = document.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.05


def add_text(document: Document, text: str, *, bold: bool = False,
             align=None, space_after: int | None = None) -> None:
    paragraph = document.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    if space_after is not None:
        paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(11)


def cover_letter() -> None:
    document = Document()
    configure(document)
    add_text(document, "9 September 2026", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_text(document, "Editor-in-Chief")
    add_text(document, "Internet of Things")
    add_text(document, "Elsevier")
    add_text(document, "Dear Editor,")
    add_text(
        document,
        'Please consider the manuscript "Blockchain access control for '
        'agricultural edge IoT: a 61-day Hyperledger Fabric field study" '
        'for publication in Internet of Things. The article reports a 61-day '
        'agricultural edge deployment, controlled latency and throughput '
        'campaigns, and an implementation audit grounded in released real '
        'measurements.')
    add_text(
        document,
        "The manuscript separates field observations from controlled "
        "experiments and scripted conformance tests. It uses runs as the "
        "experimental unit and reports estimates with confidence intervals. "
        "The implementation audit identifies four defects involving policy, "
        "residue recovery, signature verification and on-ledger revocation, "
        "plus one revocation-observability gap. The released corrected "
        "implementation addresses these defects with regression tests and an "
        "isolated chaincode benchmark. Historical performance is not "
        "presented as corrected-code performance.")
    add_text(
        document,
        "The supporting data and code are publicly archived in Zenodo at "
        "https://doi.org/10.5281/zenodo.22667556. All authors have approved "
        "the manuscript, and the work is not under consideration elsewhere.")
    add_text(document, "Sincerely,", space_after=18)
    add_text(document, "Forcha Glen Beloa", bold=True, space_after=0)
    add_text(document, "Corresponding author", space_after=0)
    add_text(document, "glenbeloa@gmail.com", space_after=0)
    document.core_properties.title = "Cover letter for Internet of Things"
    document.core_properties.author = "Forcha Glen Beloa"
    document.core_properties.subject = (
        "Blockchain access control for agricultural edge IoT")
    document.save(HERE / "Cover_letter.docx")


def highlights() -> None:
    lines = [
        "A 61-day farm deployment recorded 209,000 access decisions",
        "Counterbalanced tests linked more configured peers with lower write latency",
        "Access control reduced tested throughput by 9.6% on average",
        "A code audit exposed policy, residue, signature and revocation defects",
        "Released data and code support independent reproduction",
    ]
    if not 3 <= len(lines) <= 5 or any(len(line) > 85 for line in lines):
        raise ValueError("Highlights must contain 3 to 5 lines of at most 85 characters")
    document = Document()
    configure(document)
    add_text(document, "Highlights", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for line in lines:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(line)
        run.font.name = FONT
        run.font.size = Pt(11)
    document.core_properties.title = "Highlights"
    document.core_properties.author = "Forcha Glen Beloa and co-authors"
    document.save(HERE / "Highlights.docx")


if __name__ == "__main__":
    cover_letter()
    highlights()
